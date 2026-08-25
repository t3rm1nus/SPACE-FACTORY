"""Verificación E2E REAL: libro languages='en' a través del Autopilot completo.

Ejecuta el pipeline editorial REAL (create_job + run_job con el ejecutor de
producción: load_modules() + capabilities_map + scheduler/tasks reales, SIN
stubs) para un libro con languages='en', y verifica:

  1. job status == COMPLETED
  2. chapters.draft_en / edited_en poblados (sin placeholders)
  3. output/docx/book_{id}_en.docx existe
  4. El DOCX contiene UI en inglés ("Table of Contents", no "Índice") y texto
     de capítulo en inglés real.

Modo determinista sin LLM (mismo patrón que run_e2e_001_editorial.py):
  - CHAP_USE_LLM=0  -> Chapter Writer 100% Python (backstop determinista EN).
  - CHAP_FORCE_MIN=1 -> garantiza el mínimo operativo de palabras.
  - OLLAMA_BASE_URL apuntado a un puerto cerrado -> los demás módulos caen a
    sus fallbacks deterministas RÁPIDO (connection refused inmediato).
Research sigue siendo REAL (HTTP a Wikipedia/SearXNG): si la red no está
disponible, el gate de research fallará y el script lo reportará como tal.

Uso:
    python tools/dev/archive/verify_e2e_en_book.py
"""

from __future__ import annotations

import json
import os
import sys
import tempfile
import traceback

# Ejecutable desde cualquier cwd: añade la raíz del proyecto al path.
_PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
if _PROJECT_ROOT not in sys.path:
    sys.path.insert(0, _PROJECT_ROOT)
os.chdir(_PROJECT_ROOT)

# ---------------------------------------------------------------------------
# Entorno determinista ANTES de cualquier import de producción.
# ---------------------------------------------------------------------------
os.environ["CHAP_USE_LLM"] = "0"
os.environ["CHAP_FORCE_MIN"] = "1"
os.environ["LLM_PROVIDER"] = "ollama"
os.environ["OLLAMA_BASE_URL"] = "http://127.0.0.1:9"  # puerto cerrado: fallback rápido
os.environ["ROUTER_MODEL"] = "qwen-agent:latest"

from core.autopilot import BookJobStore, create_job, default_executor_factory, run_job
from core.database import init_db
from core.module_registry import capabilities_map, load_modules
from frontend.editorial import create_book, get_chapters

BOOK_TITLE = "The history of the printing press"
IDEA = ("The history of the printing press: from Gutenberg's movable type "
        "to industrial printing and modern digital presses.")


def log(msg: str) -> None:
    print(msg, flush=True)


def main() -> int:
    # BD aislada (nunca tocar producción desde este verificador).
    tmp_db = tempfile.NamedTemporaryFile(delete=False, suffix=".db")
    tmp_db.close()
    os.environ["SPACE_LAIR_DB_PATH"] = tmp_db.name
    init_db()

    jobs_dir = os.path.join(tempfile.mkdtemp(prefix="e2e_en_jobs_"))
    store = BookJobStore(jobs_dir)

    # --- Libro EN real -----------------------------------------------------
    created = create_book({
        "title": BOOK_TITLE,
        "idea": IDEA,
        "description": IDEA,
        "author": "Space Lair E2E",
        "genre": "Technology",
        "target_audience": "General",
        "language": "en",
        "target_chapters": 1,
        "image_count": 0,
        "image_search_ratio": 0.0,
    })
    book_id = created["book_id"]
    log(f"[setup] book_id={book_id} language=en target_chapters=1 image_count=0")

    # --- Ejecutor de PRODUCCIÓN (módulos reales, sin stubs) -----------------
    modules = load_modules()
    cap_map = capabilities_map(modules)
    executor = default_executor_factory(modules, cap_map, store)
    log(f"[setup] modules={len(modules)} capabilities={len(cap_map)}")

    # --- Job real del Autopilot --------------------------------------------
    job = create_job(store, book_id, data={"num_images": 0})
    log(f"[job] {job['job_id']} fases={[p['id'] for p in job['phases']]}")

    def emit(event: str, data: dict) -> None:
        if event in ("phase_started", "phase_completed", "phase_failed",
                     "job_started", "job_completed", "job_failed"):
            extra = ""
            if event == "phase_failed":
                extra = f" error={data.get('error')}"
            if event == "phase_completed":
                extra = f" duration={data.get('duration')}"
            log(f"[{event}] {data.get('phase') or ''}{extra}")

    final = run_job(job, store, executor, emit=emit)
    return _verify(final)



def _verify(final: dict) -> int:
    """Chequeos 1-4 sobre el job finalizado. Devuelve 0 (PASS) o 1 (FAIL)."""
    report = {"job_status": final.get("status"), "checks": {}}
    phases_state = {
        p["id"]: (p.get("status"), p.get("error"))
        for p in final.get("phases", [])
    }
    log("[phases] " + json.dumps(
        {k: v[0] for k, v in phases_state.items()}, ensure_ascii=False))
    for pid, (st, err) in phases_state.items():
        if st != "PASS":
            log(f"[phase-detail] {pid}: status={st} error={err}")
    report["phases"] = {k: v[0] for k, v in phases_state.items()}
    report["phase_errors"] = {k: v[1] for k, v in phases_state.items() if v[1]}
    checks = report["checks"]

    # Check 1: job completado
    checks["job_completed"] = final.get("status") == "COMPLETED"

    # Check 2: draft_en / edited_en poblados, sin placeholders, en inglés
    from frontend.editorial import get_chapters as _get_chs
    chs = _get_chs(final["book_id"])
    c1 = chs[0]
    draft_en = (c1.get("draft_en") or "").strip()
    edited_en = (c1.get("edited_en") or "").strip()
    words = len(draft_en.split()) if draft_en else 0
    bad_markers = ["Desarrollar el n", "contenido de prueba", "texto de ejemplo",
                   "{{", "[pendiente]", "insert text"]
    has_placeholder = any(m.lower() in draft_en.lower() for m in bad_markers)
    spanish_markers = ("el capítulo", "en este apartado", "se trata de",
                       "cabe precisar")
    looks_spanish = any(m in draft_en[:600].lower() for m in spanish_markers)
    checks["draft_en_populated"] = bool(draft_en) and words >= 1500
    checks["edited_en_populated"] = bool(edited_en)
    checks["draft_en_no_placeholder"] = bool(draft_en) and not has_placeholder
    checks["draft_en_looks_english"] = bool(draft_en) and not looks_spanish
    log(f"[db] draft_en words={words} edited_len={len(edited_en)} "
        f"placeholder={has_placeholder} looks_spanish={looks_spanish}")
    log(f"[db] draft_en head: {draft_en[:220]!r}")

    # Check 3: DOCX EN existe
    docx_path = final.get("docx_path")
    checks["docx_path_reported"] = bool(docx_path)
    checks["docx_exists"] = bool(docx_path) and os.path.isfile(docx_path)
    log(f"[docx] path={docx_path} exists={checks['docx_exists']}")

    # Check 4: contenido del DOCX en inglés
    if checks["docx_exists"]:
        from docx import Document as DocxDocument
        doc = DocxDocument(docx_path)
        texts = [p.text for p in doc.paragraphs]
        checks["toc_english"] = "Table of Contents" in texts
        checks["toc_not_spanish"] = not any(t == "Índice" for t in texts)
        checks["body_english_present"] = any(
            ("printing press" in t.lower()) or ("gutenberg" in t.lower())
            for t in texts)
        checks["doc_not_empty"] = len([t for t in texts if t.strip()]) > 10
        log(f"[docx] toc_english={checks['toc_english']} "
            f"body_english={checks['body_english_present']} "
            f"paragraphs={len(texts)}")
    else:
        checks.update({"toc_english": False, "toc_not_spanish": False,
                       "body_english_present": False, "doc_not_empty": False})

    ok = all([
        checks["job_completed"],
        checks["draft_en_populated"],
        checks["edited_en_populated"],
        checks["draft_en_no_placeholder"],
        checks["draft_en_looks_english"],
        checks["docx_exists"],
        checks["toc_english"],
        checks["toc_not_spanish"],
        checks["body_english_present"],
        checks["doc_not_empty"],
    ])
    report["VERDICT"] = "PASS" if ok else "FAIL"
    log("[report] " + json.dumps(report, ensure_ascii=False, indent=2, default=str))
    return 0 if ok else 1


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception:
        traceback.print_exc()
        sys.exit(2)
